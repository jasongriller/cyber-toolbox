"""Ingest-path tests for legacy .doc uploads."""

from __future__ import annotations

import io
import json
import re
import zipfile
from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from rmf_migrator.common.http import HttpError
from rmf_migrator.common.limits import (
    ConversionMissing,
    sniff_document_format,
)
from rmf_migrator.common.models import (
    Document,
    DocumentStatus,
    Draft,
    DraftStatus,
    ExportJob,
    JobStatus,
    ParseJob,
    Project,
)
from rmf_migrator.common.storage import (
    DOC_CONTENT_TYPE,
    DOCX_CONTENT_TYPE,
    build_document_key,
)
from rmf_migrator.doc_convert import (
    CONVERSION_DISABLED_MESSAGE,
    ConversionFailed,
    ConversionUnavailable,
)
from rmf_migrator.docx.parser import parse_docx_bytes
from rmf_migrator.handlers.parse_document import run_parse_job
from rmf_migrator.handlers.request_upload import _request_upload
from rmf_migrator.handlers.worker import run_export_job
from tests.conftest import ole2_container

FIXTURES = Path(__file__).parent / "fixtures"

# `_ensure_docx` routes on sniff_document_format, and since Task 1's amendment
# that means a real Word directory entry — bare OLE2 magic classifies as None.
OLE2 = ole2_container("WordDocument", "1Table")


def _event(project_id: str, filename: str) -> dict:
    return {
        "pathParameters": {"project_id": project_id},
        "body": json.dumps({"filename": filename}),
    }


def test_doc_upload_rejected_when_conversion_disabled(deps_with_project):
    deps, project_id = deps_with_project
    deps.config = replace(deps.config, doc_conversion_backend="reject")

    with pytest.raises(HttpError) as exc:
        _request_upload(_event(project_id, "policy.doc"), deps)

    assert exc.value.status == 400
    # Equality, not a substring: this is what proves the handler imports the
    # constant rather than retyping the text and letting the two drift.
    assert exc.value.message == CONVERSION_DISABLED_MESSAGE
    # The refusal has to land before anything is written. A guard that ran
    # after registration would leave a Document pointing at an S3 key nothing
    # will ever upload, and would inflate the project's count for good.
    assert deps.repo.get_project(project_id).document_count == 0
    assert deps.repo.list_documents(project_id) == []


def test_doc_upload_accepted_when_conversion_enabled(deps_with_conversion, capsys):
    deps, project_id = deps_with_conversion

    response = _request_upload(_event(project_id, "policy.doc"), deps)

    assert response["statusCode"] == 201
    body = json.loads(response["body"])
    assert body["document"]["source_format"] == "doc"
    assert body["document"]["s3_key"].endswith(".doc")
    assert body["upload"]["fields"]["Content-Type"] == DOC_CONTENT_TYPE
    # The audit trail is the only durable record of who exercised the converter:
    # the Document row carries source_format too, but it is application state
    # that dies with the project. Without this field a legacy registration is
    # byte-identical to a .docx one, so no metric filter or alarm can count the
    # legacy path or alert on a rate spike against it.
    events = [json.loads(line) for line in capsys.readouterr().out.strip().splitlines()]
    assert [event["event"] for event in events] == ["document.registered"]
    assert events[0]["source_format"] == "doc"


def test_docx_upload_is_unaffected_by_conversion_config(deps_with_conversion):
    deps, project_id = deps_with_conversion

    response = _request_upload(_event(project_id, "policy.docx"), deps)

    assert response["statusCode"] == 201
    assert json.loads(response["body"])["document"]["source_format"] == "docx"


def test_other_extensions_still_rejected(deps_with_conversion):
    deps, project_id = deps_with_conversion

    with pytest.raises(HttpError) as exc:
        _request_upload(_event(project_id, "policy.pdf"), deps)

    assert exc.value.status == 400


# --- worker: convert, then parse --------------------------------------------


def _styled_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Access Control Policy", level=1)
    doc.add_paragraph("The organization manages accounts.")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _truncated_zip_bytes() -> bytes:
    """Zip magic over a cut-off archive — what a converter that died mid-write emits."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("word/document.xml", "<w:document/>")
    return out.getvalue()[:32]


def _seed_job(deps, filename: str, data: bytes, source_format: str) -> tuple[str, str, str]:
    """Seed a project, an UPLOADED document with its bytes in S3, and a PENDING job."""
    project = Project(name="Sys")
    deps.repo.put_project(project)
    document = Document(
        project_id=project.project_id,
        filename=filename,
        s3_key="",  # set below once we have the id, as request_upload does
        source_format=source_format,
    )
    document.s3_key = build_document_key(project.project_id, document.document_id, filename)
    document.status = DocumentStatus.UPLOADED
    deps.repo.put_document(document)
    deps.store.put_bytes(
        document.s3_key,
        data,
        DOC_CONTENT_TYPE if source_format == "doc" else DOCX_CONTENT_TYPE,
    )
    job = ParseJob(project_id=project.project_id, document_id=document.document_id)
    deps.repo.put_job(job)
    return project.project_id, document.document_id, job.job_id


@pytest.fixture
def doc_job(deps_with_conversion):
    deps, _ = deps_with_conversion
    return (deps, *_seed_job(deps, "policy.doc", OLE2, "doc"))


@pytest.fixture
def docx_job(deps_with_conversion):
    deps, _ = deps_with_conversion
    return (deps, *_seed_job(deps, "policy.docx", _styled_docx_bytes(), "docx"))


class _StubConverter:
    def __init__(self, output: bytes | None = None, error: Exception | None = None):
        self.output = output
        self.error = error
        self.calls: list[bytes] = []

    def convert(self, data: bytes) -> bytes:
        self.calls.append(data)
        if self.error:
            raise self.error
        return self.output


def test_doc_document_is_converted_then_parsed(doc_job):
    deps, project_id, document_id, job_id = doc_job
    deps.converter = _StubConverter(output=_styled_docx_bytes())

    assert run_parse_job(project_id, document_id, job_id, deps) is True

    document = deps.repo.get_document(project_id, document_id)
    assert document.status == DocumentStatus.PARSED
    assert document.section_count > 0
    assert deps.converter.calls == [OLE2]


def test_converted_bytes_are_stored_under_a_separate_key(doc_job):
    deps, project_id, document_id, job_id = doc_job
    deps.converter = _StubConverter(output=_styled_docx_bytes())
    run_parse_job(project_id, document_id, job_id, deps)

    document = deps.repo.get_document(project_id, document_id)
    assert document.converted_s3_key is not None
    assert document.converted_s3_key != document.s3_key
    assert document.parse_key() == document.converted_s3_key
    # Original .doc preserved byte-for-byte as audit provenance.
    assert deps.store.get_bytes(document.s3_key, max_bytes=None) == OLE2
    # The stored object, not just the in-memory value the parser saw. Parsing
    # cannot catch a bad write here — it reads the converter's return value
    # directly — so without this the document reaches PARSED, survives human
    # review, and only dies at export, where parse_key() is finally fetched.
    assert deps.store.get_bytes(document.converted_s3_key) == deps.converter.output


def test_docx_document_never_touches_the_converter(docx_job):
    deps, project_id, document_id, job_id = docx_job
    deps.converter = _StubConverter(error=AssertionError("must not be called"))

    assert run_parse_job(project_id, document_id, job_id, deps) is True
    assert deps.converter.calls == []


def test_conversion_failure_marks_the_document_failed_at_the_convert_stage(doc_job):
    deps, project_id, document_id, job_id = doc_job
    deps.converter = _StubConverter(error=ConversionFailed("soffice died"))

    with pytest.raises(ConversionFailed):
        run_parse_job(project_id, document_id, job_id, deps)

    document = deps.repo.get_document(project_id, document_id)
    assert document.status == DocumentStatus.FAILED
    assert document.failure_stage == "convert"
    assert document.parse_error == "ConversionFailed"
    # Error *type* only — no converter message leaks into the record.
    assert "soffice" not in (document.parse_error or "")

    job = deps.repo.get_job(project_id, job_id)
    assert job.status == JobStatus.FAILED
    assert job.error_type == "ConversionFailed"


def test_a_convert_failure_is_retryable_by_a_redelivered_message(doc_job):
    deps, project_id, document_id, job_id = doc_job
    deps.converter = _StubConverter(error=ConversionUnavailable("backend off"))
    with pytest.raises(ConversionUnavailable):
        run_parse_job(project_id, document_id, job_id, deps)

    # ConversionUnavailable is a deployment state: the operator enables the
    # backend and the same SQS message is redelivered. The re-parse gate must
    # re-admit a convert failure as it re-admits a parse failure, or the fix
    # can never take effect and the retry dies as "InvalidDocumentState".
    deps.converter = _StubConverter(output=_styled_docx_bytes())
    assert run_parse_job(project_id, document_id, job_id, deps) is True

    document = deps.repo.get_document(project_id, document_id)
    assert document.status == DocumentStatus.PARSED
    assert document.failure_stage is None
    job = deps.repo.get_job(project_id, job_id)
    assert job.status == JobStatus.SUCCEEDED


def test_a_downstream_failure_is_still_refused_by_the_re_parse_gate(doc_job):
    deps, project_id, document_id, job_id = doc_job
    document = deps.repo.get_document(project_id, document_id)
    document.status = DocumentStatus.FAILED
    document.failure_stage = "mapping"
    deps.repo.put_document(document)
    deps.converter = _StubConverter(error=AssertionError("must not be called"))

    # The other half of the widened gate. Convert is upstream of parse, so it
    # re-admits; mapping is downstream, and a mapping job may still be in
    # flight over sections a re-parse would rewrite underneath it.
    assert run_parse_job(project_id, document_id, job_id, deps) is False
    job = deps.repo.get_job(project_id, job_id)
    assert job.error_type == "InvalidDocumentState"


def test_doc_bytes_under_a_docx_name_are_recorded_as_legacy(deps_with_conversion):
    deps, _ = deps_with_conversion
    # The mirror image of the case below: a legacy .doc sent as "policy.docx"
    # is on record as native, but the worker routes on bytes and converts it.
    # Leaving source_format "docx" would put a converted document on record as
    # native, and the UI's conversion badge reads that field — reviewers would
    # be told the formatting is the author's when it came out of a converter.
    project_id, document_id, job_id = _seed_job(deps, "policy.docx", OLE2, "docx")
    deps.converter = _StubConverter(output=_styled_docx_bytes())

    assert run_parse_job(project_id, document_id, job_id, deps) is True

    document = deps.repo.get_document(project_id, document_id)
    assert document.status == DocumentStatus.PARSED
    assert document.converted_s3_key is not None
    assert document.source_format == "doc"


def test_a_docx_misnamed_doc_parses_and_then_exports(deps_with_conversion):
    deps, _ = deps_with_conversion
    # request_upload derives source_format from the extension, so a genuine
    # .docx sent as "policy.doc" is on record as legacy. The worker routes on
    # bytes and converts nothing, which leaves no converted key — provenance
    # has to follow that content decision, or the document parses cleanly and
    # is then un-exportable forever because parse_key refuses it.
    project_id, document_id, job_id = _seed_job(deps, "policy.doc", _styled_docx_bytes(), "doc")
    deps.converter = _StubConverter(error=AssertionError("must not be called"))

    assert run_parse_job(project_id, document_id, job_id, deps) is True
    document = deps.repo.get_document(project_id, document_id)
    assert document.status == DocumentStatus.PARSED
    assert document.converted_s3_key is None
    assert document.parse_key() == document.s3_key

    sections = deps.repo.list_sections(document_id)
    deps.repo.put_drafts(
        [
            Draft(
                project_id=project_id,
                document_id=document_id,
                section_id=section.section_id,
                order=section.order,
                draft_text="Rev 5 text.",
                status=DraftStatus.APPROVED,
            )
            for section in sections
        ]
    )
    document.status = DocumentStatus.EXPORTING
    deps.repo.put_document(document)
    export_job = ExportJob(project_id=project_id, document_id=document_id)
    deps.repo.put_export_job(export_job)

    run_export_job(project_id, document_id, export_job.job_id, deps)

    document = deps.repo.get_document(project_id, document_id)
    assert document.status == DocumentStatus.EXPORTED
    assert deps.repo.get_export_job(project_id, export_job.job_id).status == JobStatus.SUCCEEDED


def test_a_re_parse_of_replaced_bytes_retires_the_stale_converted_key(deps_with_conversion):
    deps, _ = deps_with_conversion
    project_id, document_id, job_id = _seed_job(deps, "policy.doc", OLE2, "doc")
    # Conversion lands its key on the record and the parse of that output then
    # fails, leaving the document FAILED with a converted key on file.
    deps.converter = _StubConverter(output=_truncated_zip_bytes())
    # UnsupportedDocumentFormat is a permanent verdict on the bytes (retrying
    # the same content changes nothing), so run_parse_job records it and ends
    # the message rather than propagating it for SQS to redeliver.
    assert run_parse_job(project_id, document_id, job_id, deps) is False
    document = deps.repo.get_document(project_id, document_id)
    assert document.converted_s3_key is not None

    # Still inside the upload POST's TTL the uploader can overwrite s3_key, and
    # a FAILED parse is re-admitted. The whole provenance pair has to follow the
    # new content decision, not just source_format: parse_key prefers the
    # converted key whenever it is set, so a leftover one would hand the
    # exporter bytes that were never parsed, mapped, or reviewed.
    deps.store.put_bytes(document.s3_key, _styled_docx_bytes())
    deps.converter = _StubConverter(error=AssertionError("must not be called"))

    assert run_parse_job(project_id, document_id, job_id, deps) is True

    document = deps.repo.get_document(project_id, document_id)
    assert document.source_format == "docx"
    assert document.converted_s3_key is None
    assert document.parse_key() == document.s3_key


def test_a_failed_conversion_leaves_the_record_unparseable_rather_than_native(
    deps_with_conversion,
):
    deps, _ = deps_with_conversion
    project_id, document_id, job_id = _seed_job(deps, "policy.docx", OLE2, "docx")
    deps.converter = _StubConverter(error=ConversionFailed("soffice died"))

    with pytest.raises(ConversionFailed):
        run_parse_job(project_id, document_id, job_id, deps)

    # The record is corrected to "doc" *before* the converter runs, so a failed
    # conversion strands the document instead of leaving it on record as native
    # — which would let raw OLE2 bytes resolve through parse_key and reach the
    # exporter. Moving that assignment down next to the conversion silently
    # reintroduces exactly that hazard, so pin the ordering here.
    document = deps.repo.get_document(project_id, document_id)
    assert document.source_format == "doc"
    assert document.converted_s3_key is None
    with pytest.raises(ConversionMissing):
        document.parse_key()


def test_a_failed_conversion_retires_an_earlier_attempts_converted_key(deps_with_conversion):
    deps, _ = deps_with_conversion
    project_id, document_id, job_id = _seed_job(deps, "policy.doc", OLE2, "doc")
    # First attempt converts, lands its key on the record, and then fails to
    # parse the converter's output — FAILED at "parse" with a converted key.
    deps.converter = _StubConverter(output=_truncated_zip_bytes())
    # UnsupportedDocumentFormat is a permanent verdict on the bytes, so
    # run_parse_job records it and ends the message rather than propagating.
    assert run_parse_job(project_id, document_id, job_id, deps) is False
    stale_key = deps.repo.get_document(project_id, document_id).converted_s3_key
    assert stale_key is not None

    # The uploader overwrites s3_key with different legacy bytes inside the
    # POST's TTL and the FAILED document is re-admitted, but this conversion
    # fails. build_converted_document_key is deterministic, so the key value is
    # unchanged while the S3 object behind it is the *first* attempt's output —
    # content nobody parsed and that is no longer at s3_key. The key has to
    # retire before the converter runs, exactly as source_format is corrected
    # before it, or parse_key keeps resolving that stale object for the exporter.
    document = deps.repo.get_document(project_id, document_id)
    deps.store.put_bytes(
        document.s3_key, ole2_container("WordDocument", "1Table", "SummaryInformation")
    )
    deps.converter = _StubConverter(error=ConversionFailed("soffice died"))

    with pytest.raises(ConversionFailed):
        run_parse_job(project_id, document_id, job_id, deps)

    document = deps.repo.get_document(project_id, document_id)
    assert document.converted_s3_key is None
    with pytest.raises(ConversionMissing):
        document.parse_key()


def test_a_corrected_format_is_auditable_from_the_logs(deps_with_conversion, capsys):
    deps, _ = deps_with_conversion
    project_id, document_id, job_id = _seed_job(deps, "policy.docx", OLE2, "docx")
    deps.converter = _StubConverter(output=_styled_docx_bytes())
    capsys.readouterr()

    assert run_parse_job(project_id, document_id, job_id, deps) is True

    # document.registered counted this upload as native, so this population of
    # conversions never appears in the registration count. Without a signal at
    # the flip the only record is the DynamoDB row — application state that
    # dies with the project, and that no metric filter or alarm can read.
    events = [json.loads(line) for line in capsys.readouterr().out.strip().splitlines()]
    corrected = [event for event in events if event["event"] == "document.format_corrected"]
    assert len(corrected) == 1
    assert corrected[0]["document_id"] == document_id
    assert corrected[0]["registered_format"] == "docx"
    assert corrected[0]["sniffed_format"] == "doc"


def test_the_committed_doc_fixture_is_genuinely_legacy_binary():
    """Keeps the fidelity test below from going vacuous. Nothing else would
    notice if legacy_policy.doc were replaced by a renamed .docx, and the pair
    would then demonstrate nothing at all about LibreOffice."""
    assert sniff_document_format((FIXTURES / "legacy_policy.doc").read_bytes()) == "doc"


def _pinned_libreoffice_series() -> str:
    """The LibreOffice series converter.Dockerfile refuses to build without.

    Read from the Dockerfile rather than restated, so that the fixture's
    provenance is pinned to the converter this repo actually ships. Bumping the
    base image and leaving the fixture behind is the likelier mistake than
    swapping the fixture, and it fails open: the evidence would quietly stop
    describing the shipped converter.
    """
    match = re.search(
        r"grep -q 'LibreOffice (\d+)\\\.(\d+)\\\.'",
        (Path(__file__).parent.parent / "converter.Dockerfile").read_text(),
    )
    assert match, "converter.Dockerfile no longer pins a LibreOffice series at build time"
    return f"{match.group(1)}.{match.group(2)}."


def test_the_committed_converted_fixture_came_from_the_pinned_converter():
    """The other half of that guard, for the other half of the pair. Without it
    the .docx can be swapped for a hand-authored python-docx file and every
    assertion below still passes, measuring nothing about LibreOffice. Reading
    the expected series out of converter.Dockerfile also breaks this test on a
    base-image bump that leaves the fixture behind, instead of letting the
    fixture quietly become evidence about a converter nobody ships."""
    app_xml = (
        zipfile.ZipFile(io.BytesIO((FIXTURES / "legacy_policy.converted.docx").read_bytes()))
        .read("docProps/app.xml")
        .decode()
    )
    # LibreOffice writes "LibreOffice/7.4.7.2$Linux_X86_64 ...", python-docx
    # writes "Microsoft Macintosh Word", Word writes "Microsoft Office Word".
    series = _pinned_libreoffice_series()
    assert f"LibreOffice/{series}" in app_xml, (
        f"fixture was not produced by the converter image pinned to "
        f"LibreOffice {series}x: {app_xml}"
    )


def test_conversion_preserves_heading_structure():
    """The whole feature rests on this: if conversion loses heading styles,
    every document collapses into one unsectioned blob and the mapping
    workflow silently degrades.

    The .docx read here is not authored — it is what the converter image
    produced from the committed .doc, so the assertions land on a real
    LibreOffice conversion rather than a stand-in that would pass regardless.
    """
    converted = (FIXTURES / "legacy_policy.converted.docx").read_bytes()
    assert sniff_document_format(converted) == "docx"

    sections = parse_docx_bytes(converted, document_id="doc_1", project_id="proj_1")
    headings = {section.heading: section for section in sections if section.heading}

    # Named up front so a total loss of heading styles reports as itself rather
    # than as a KeyError three lines down.
    assert set(headings) == {"Access Control Policy", "Account Management"}, (
        f"headings did not survive conversion: {sorted(headings)}"
    )
    assert headings["Access Control Policy"].level == 1, "top-level heading did not survive"
    assert headings["Account Management"].level == 2, "nested heading did not survive"
    # Depth alone is not enough — mapping walks parents to scope a control.
    assert headings["Account Management"].parent_id == headings["Access Control Policy"].section_id
    # Body text under each heading, not merely somewhere in the document.
    assert "manages information system accounts" in headings["Access Control Policy"].text
    assert "reviewed annually" in headings["Account Management"].text
